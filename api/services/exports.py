"""Portfolio exports (Phase 5) — CSV, master JSON, and XLSX.

Every export reflects DB/view-layer data exactly (built from the same portfolio and
view models the API serves — no recomputation). No ESPN traffic; no secrets.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import UTC, datetime

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import DraftPick, League, Matchup, Team, Transaction
from ..schemas import DraftPickOut, MatchupOut, TeamOut, TransactionOut
from .draft_analytics import build_draft_adp, build_strategies
from .exposure import ExposureScope, ExposureView, build_exposure, players_for_view
from .metrics import team_edge
from .portfolio import build_portfolio_rows, build_summary
from .portfolio_filters import PortfolioFilters
from .read_models import build_league_out

# Column order shared by CSV + XLSX so the two never drift. Edge Index (Phase 16/17) is the
# primary advantage score; legacy edge_score/grade/verdict are kept for compatibility.
_ROW_FIELDS = [
    "league_id", "season", "league_name", "size", "account_label", "lifecycle",
    "my_team_name", "wins", "losses", "ties", "points_for", "points_against",
    "standing", "edge_index_score", "edge_index_grade", "edge_index_verdict",
    "edge_score", "grade", "verdict", "playoff_odds", "last_synced_at",
]

_EXPOSURE_ROW_FIELDS = [
    "espn_player_id", "player_name", "position", "nfl_team", "rostered_teams",
    "teams_in_scope", "exposure_pct", "share", "rostered_leagues",
    "leagues_in_scope", "league_exposure_pct", "league_share",
    "my_rostered_teams", "my_teams_in_scope", "my_rostered_leagues",
    "my_leagues_in_scope", "my_exposure_pct", "my_share",
    "field_rostered_teams", "field_teams_in_scope", "field_rostered_leagues",
    "field_leagues_in_scope", "field_exposure_pct", "field_share",
    "field_slot_pct", "field_slot_share", "leverage_pp", "avg_overall",
    "min_overall", "max_overall", "avg_pick_value", "auction_rosters",
    "leagues",
]
_EXPOSURE_HEADLINE_ROW_FIELDS = [
    "headline", "espn_player_id", "player_name", "position", "nfl_team",
    "exposure_pct", "field_exposure_pct", "leverage_pp", "share", "field_share",
    "field_slot_pct", "field_slot_share",
    "pick_value_pct", "field_pick_value_pct", "pick_count", "field_pick_count",
    "penetration_pct", "penetration_share", "players_per_team",
    "players_per_team_share", "draft_time_adp", "current_ffc_adp",
    "market_move", "market_move_abs", "market_move_label",
]
_NFL_CONCENTRATION_ROW_FIELDS = [
    "nfl_team", "teams_with_player", "teams_in_scope", "penetration_pct",
    "penetration_share", "player_team_instances", "players_per_team",
    "players_per_team_share",
]
_POSITIONAL_SPEND_ROW_FIELDS = [
    "position", "pick_count", "pick_value", "pick_value_pct", "total_pick_value",
    "field_pick_count", "field_pick_value", "field_pick_value_pct",
    "field_total_pick_value", "leverage_pp",
]
_DRAFT_ADP_ROW_FIELDS = [
    "league_id", "league_name", "team_id", "team_name", "draft_type",
    "draft_value_capture_espn", "draft_value_capture_ffc",
    "draft_adp_source_disagreement",
    "draft_value_capture_espn_portfolio_median",
    "draft_value_capture_espn_vs_portfolio_median",
    "draft_value_capture_ffc_portfolio_median",
    "draft_value_capture_ffc_vs_portfolio_median",
    "draft_adp_source_disagreement_portfolio_median",
    "draft_adp_source_disagreement_vs_portfolio_median",
]
_STRATEGY_ROW_FIELDS = [
    "league_id", "league_name", "team_id", "team_name", "draft_type",
    "primary_label", "primary_confidence", "secondary_label",
    "secondary_confidence", "edge_index_score", "triggering_picks",
]


def _teams_sorted(session: Session, league_id: int) -> list[Team]:
    teams = session.scalars(select(Team).where(Team.league_id == league_id))
    return sorted(teams, key=lambda t: (t.standing is None, t.standing or 0))


# --------------------------------------------------------------------------- #
# CSV — one row per league (the Portfolio Board columns)
# --------------------------------------------------------------------------- #
def portfolio_csv(session: Session) -> str:
    rows = build_portfolio_rows(session)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=_ROW_FIELDS, extrasaction="ignore")
    writer.writeheader()
    for r in rows:
        writer.writerow({k: getattr(r, k) for k in _ROW_FIELDS})
    return buf.getvalue()


def _csv_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, separators=(",", ":"), sort_keys=True)
    return value


def _dict_rows_csv(rows: list[dict], fields: list[str]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({field: _csv_value(row.get(field)) for field in fields})
    return buf.getvalue()


def exposure_csv(
    session: Session,
    *,
    scope: ExposureScope,
    view: ExposureView,
    filters: PortfolioFilters,
) -> str:
    data = build_exposure(session, scope=scope, filters=filters)
    return _dict_rows_csv(players_for_view(data["players"], view), _EXPOSURE_ROW_FIELDS)


def draft_adp_csv(session: Session, *, filters: PortfolioFilters) -> str:
    data = build_draft_adp(session, filters)
    return _dict_rows_csv(data["teams"], _DRAFT_ADP_ROW_FIELDS)


def strategies_csv(session: Session, *, filters: PortfolioFilters) -> str:
    data = build_strategies(session, filters)
    return _dict_rows_csv(data["teams"], _STRATEGY_ROW_FIELDS)


# --------------------------------------------------------------------------- #
# Master JSON — summary + rows + per-league detail (mirrors the API models)
# --------------------------------------------------------------------------- #
def portfolio_json(session: Session) -> dict:
    rows = build_portfolio_rows(session)
    summary = build_summary(rows)
    leagues_out = []
    for lg in session.scalars(select(League).order_by(League.season.desc(), League.id)):
        teams = _teams_sorted(session, lg.id)
        draft = session.scalars(
            select(DraftPick).where(DraftPick.league_id == lg.id).order_by(DraftPick.overall)
        )
        matchups = session.scalars(
            select(Matchup).where(Matchup.league_id == lg.id).order_by(Matchup.week)
        )
        activity = session.scalars(select(Transaction).where(Transaction.league_id == lg.id))
        leagues_out.append(
            {
                "league": build_league_out(session, lg).model_dump(mode="json"),
                "teams": [TeamOut.model_validate(t).model_dump(mode="json") for t in teams],
                "draft": [DraftPickOut.model_validate(p).model_dump(mode="json") for p in draft],
                "matchups": [
                    MatchupOut.model_validate(m).model_dump(mode="json") for m in matchups
                ],
                "activity": [
                    TransactionOut.model_validate(x).model_dump(mode="json") for x in activity
                ],
            }
        )
    filters = PortfolioFilters()
    return {
        "generated_at": datetime.now(UTC).isoformat(),
        "summary": summary.model_dump(mode="json"),
        "rows": [r.model_dump(mode="json") for r in rows],
        "leagues": leagues_out,
        "exposure": {
            "me": build_exposure(session, scope="me", filters=filters),
            "opponents": build_exposure(session, scope="opponents", filters=filters),
        },
        "draft_adp": build_draft_adp(session, filters),
        "strategies": build_strategies(session, filters),
    }


# --------------------------------------------------------------------------- #
# XLSX — Portfolio summary sheet + one sheet per league
# --------------------------------------------------------------------------- #
_INVALID_SHEET = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(name: str, used: set[str]) -> str:
    title = _INVALID_SHEET.sub("-", name or "League")[:28].strip() or "League"
    candidate = title
    i = 2
    while candidate.lower() in used:
        candidate = f"{title[:26]}-{i}"
        i += 1
    used.add(candidate.lower())
    return candidate


def _append_dict_sheet(
    wb: Workbook,
    used: set[str],
    title: str,
    rows: list[dict],
    fields: list[str],
) -> None:
    sheet = wb.create_sheet(_sheet_title(title, used))
    sheet.append([field.replace("_", " ").title() for field in fields])
    for row in rows:
        sheet.append([_csv_value(row.get(field)) for field in fields])


def _exposure_headline_rows(headlines: dict) -> list[dict]:
    labels = {
        "highest_leverage": "Highest leverage",
        "most_underowned": "Most under-owned",
        "positional_capital_vs_field": "Positional capital vs field",
        "most_concentrated_nfl_team": "Most concentrated NFL team",
        "largest_market_move": "Largest market move",
    }
    rows = []
    for key, label in labels.items():
        value = headlines.get(key)
        if value is None:
            continue
        row = {"headline": label}
        row.update(value)
        rows.append(row)
    return rows


def portfolio_xlsx(session: Session) -> bytes:
    wb = Workbook()
    used: set[str] = set()

    ws = wb.active
    ws.title = _sheet_title("Portfolio", used)
    ws.append([f.replace("_", " ").title() for f in _ROW_FIELDS])
    for r in build_portfolio_rows(session):
        ws.append([getattr(r, k) for k in _ROW_FIELDS])

    filters = PortfolioFilters()
    exposure = build_exposure(session, scope="me", filters=filters)
    draft_adp = build_draft_adp(session, filters)
    strategies = build_strategies(session, filters)
    _append_dict_sheet(
        wb,
        used,
        "Exposure Rostered",
        players_for_view(exposure["players"], "rostered"),
        _EXPOSURE_ROW_FIELDS,
    )
    _append_dict_sheet(
        wb,
        used,
        "Exposure Field Owns",
        players_for_view(exposure["players"], "field_owned"),
        _EXPOSURE_ROW_FIELDS,
    )
    _append_dict_sheet(
        wb, used, "Exposure All", exposure["players"], _EXPOSURE_ROW_FIELDS
    )
    _append_dict_sheet(
        wb,
        used,
        "Exposure Headlines",
        _exposure_headline_rows(exposure["headlines"]),
        _EXPOSURE_HEADLINE_ROW_FIELDS,
    )
    _append_dict_sheet(
        wb,
        used,
        "NFL Concentration",
        exposure["nfl_team_concentration"],
        _NFL_CONCENTRATION_ROW_FIELDS,
    )
    _append_dict_sheet(
        wb,
        used,
        "Positional Spend",
        exposure["positional_spend"],
        _POSITIONAL_SPEND_ROW_FIELDS,
    )
    _append_dict_sheet(
        wb, used, "Draft ADP", draft_adp["teams"], _DRAFT_ADP_ROW_FIELDS
    )
    _append_dict_sheet(
        wb, used, "Strategies", strategies["teams"], _STRATEGY_ROW_FIELDS
    )

    header = ["Standing", "Team", "Abbrev", "W", "L", "T", "PF", "PA", "Edge", "Playoff %", "Me"]
    for lg in session.scalars(select(League).order_by(League.season.desc(), League.id)):
        sheet = wb.create_sheet(_sheet_title(f"{lg.name or lg.espn_league_id} {lg.season}", used))
        sheet.append(header)
        for t in _teams_sorted(session, lg.id):
            edge = team_edge(session, lg.id, t.id)
            sheet.append(
                [
                    t.standing,
                    t.name,
                    t.abbrev,
                    t.wins,
                    t.losses,
                    t.ties,
                    round(t.points_for, 1),
                    round(t.points_against, 1),
                    edge.edge_score,
                    None if edge.playoff_odds is None else round(edge.playoff_odds * 100, 1),
                    "yes" if t.is_me else "",
                ]
            )

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
